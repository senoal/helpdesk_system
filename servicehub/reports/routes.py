import csv
from datetime import datetime, time
from io import BytesIO, StringIO

from flask import Blueprint, Response, abort, current_app, render_template, request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import joinedload

from servicehub import db
from servicehub.models import Category, Department, Ticket, User
from servicehub.utils import local_to_utc, roles_required, utc_to_local


reports_bp = Blueprint("reports", __name__)
PRIORITIES = ["Low", "Medium", "High", "Critical"]
STATUSES = ["Open", "In Progress", "Pending", "Resolved", "Closed"]


def parse_date(value, field_name):
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        abort(400, description=f"{field_name} tidak valid.")


def report_query():
    query = db.select(Ticket).options(
        joinedload(Ticket.department), joinedload(Ticket.category),
        joinedload(Ticket.requester), joinedload(Ticket.assignee)
    )
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    parsed_from = parse_date(date_from, "Tanggal awal")
    parsed_to = parse_date(date_to, "Tanggal akhir")
    if parsed_from and parsed_to and parsed_from > parsed_to:
        abort(400, description="Tanggal awal tidak boleh setelah tanggal akhir.")
    if parsed_from:
        local_start = datetime.combine(parsed_from, time.min)
        query = query.where(Ticket.created_at >= local_to_utc(local_start, current_app.config["APP_TIMEZONE"]))
    if parsed_to:
        local_end = datetime.combine(parsed_to, time.max)
        query = query.where(Ticket.created_at <= local_to_utc(local_end, current_app.config["APP_TIMEZONE"]))
    for field, column in [("department", Ticket.department_id), ("category", Ticket.category_id), ("agent", Ticket.assignee_id)]:
        value = request.args.get(field, type=int)
        if value:
            query = query.where(column == value)
    priority = request.args.get("priority", "")
    status = request.args.get("status", "")
    if priority:
        if priority not in PRIORITIES:
            abort(400, description="Priority laporan tidak valid.")
        query = query.where(Ticket.priority == priority)
    if status:
        if status not in STATUSES:
            abort(400, description="Status laporan tidak valid.")
        query = query.where(Ticket.status == status)
    return query.order_by(Ticket.created_at.desc())


def spreadsheet_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def row_for(ticket):
    timezone_name = current_app.config["APP_TIMEZONE"]
    return [spreadsheet_safe(value) for value in [
        ticket.ticket_number, ticket.title, ticket.requester.name,
        ticket.department.name, ticket.category.name, ticket.priority,
        ticket.status, ticket.assignee.name if ticket.assignee else "",
        utc_to_local(ticket.created_at, timezone_name).strftime("%Y-%m-%d %H:%M"),
        utc_to_local(ticket.sla_deadline, timezone_name).strftime("%Y-%m-%d %H:%M"),
        utc_to_local(ticket.resolved_at, timezone_name).strftime("%Y-%m-%d %H:%M") if ticket.resolved_at else "",
        ticket.sla_state, round(ticket.resolution_hours, 2) if ticket.resolution_hours is not None else "",
    ]]


HEADERS = ["Ticket", "Judul", "Requester", "Department", "Kategori", "Priority", "Status", "Agent", "Dibuat (WIB)", "Deadline SLA (WIB)", "Selesai (WIB)", "SLA", "Waktu Penyelesaian (jam)"]


def aggregate(tickets, key_func):
    groups = {}
    for ticket in tickets:
        label = key_func(ticket) or "Unassigned"
        groups.setdefault(label, []).append(ticket)
    result = []
    for label, items in groups.items():
        completed = [ticket for ticket in items if ticket.is_completed]
        resolution = [ticket.resolution_hours for ticket in completed if ticket.resolution_hours is not None]
        met = sum(not ticket.is_overdue for ticket in completed)
        result.append({
            "label": label,
            "total": len(items),
            "open": sum(not ticket.is_completed for ticket in items),
            "completed": len(completed),
            "overdue": sum(ticket.is_overdue for ticket in items),
            "average_hours": sum(resolution) / len(resolution) if resolution else 0,
            "sla_compliance": met / len(completed) * 100 if completed else 100,
        })
    return sorted(result, key=lambda item: (-item["total"], item["label"]))


@reports_bp.route("/")
@login_required
@roles_required("Admin", "Agent")
def index():
    tickets = db.session.scalars(report_query()).unique().all()
    completed = [t for t in tickets if t.is_completed]
    resolution_times = [t.resolution_hours for t in completed if t.resolution_hours is not None]
    summary = {
        "total": len(tickets),
        "completed": len(completed),
        "average_hours": sum(resolution_times) / len(resolution_times) if resolution_times else 0,
        "sla_compliance": sum(not t.is_overdue for t in completed) / len(completed) * 100 if completed else 100,
    }
    breakdowns = {
        "Department": aggregate(tickets, lambda ticket: ticket.department.name if ticket.department else "Tanpa Department"),
        "Kategori": aggregate(tickets, lambda ticket: ticket.category.name if ticket.category else "Tanpa Kategori"),
        "Priority": aggregate(tickets, lambda ticket: ticket.priority),
        "Status": aggregate(tickets, lambda ticket: ticket.status),
        "Agent": aggregate(tickets, lambda ticket: ticket.assignee.name if ticket.assignee else "Unassigned"),
    }
    return render_template(
        "reports/index.html", tickets=tickets[:200], result_total=len(tickets), summary=summary, breakdowns=breakdowns,
        departments=db.session.scalars(db.select(Department).order_by(Department.name)).all(),
        categories=db.session.scalars(db.select(Category).order_by(Category.name)).all(),
        agents=db.session.scalars(db.select(User).where(User.role.in_(["Admin", "Agent"])).order_by(User.name)).all(),
    )


@reports_bp.route("/export.csv")
@login_required
@roles_required("Admin", "Agent")
def export_csv():
    output = StringIO()
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    for ticket in db.session.scalars(report_query()).unique():
        writer.writerow(row_for(ticket))
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=servicehub-report.csv"})


@reports_bp.route("/export.xlsx")
@login_required
@roles_required("Admin", "Agent")
def export_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ticket Report"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for ticket in db.session.scalars(report_query()).unique():
        sheet.append(row_for(ticket))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 40)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name="servicehub-report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
